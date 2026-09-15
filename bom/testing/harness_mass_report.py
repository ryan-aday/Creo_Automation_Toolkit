#!/usr/bin/env python3
"""
Creo Cabling Harness Mass Analyzer
==================================

Discovers harness routing assemblies (default token: ``_ROUT``) from a Creo
assembly mass report JSON/XLSX file, optionally exports one Creo Schematics XML
per harness through a user-recorded CREOSON mapkey, parses spool/connection
parameters, calculates wire/cable and cosmetic covering masses, and appends
harness worksheets to the existing Excel report.

Why a mapkey template?
----------------------
creopyson/CREOSON exposes ``interface_mapkey()`` but does not provide a direct
Cabling "Logical Data > Export > Creo Schematic" API. Creo menu command IDs vary
with version/site configuration, so this script deliberately replays a mapkey
recorded once in the user's own Creo installation rather than hard-coding UI
commands.

Mass conventions
----------------
PTC defines Cabling spool DENSITY as a *linear density* (mass / unit length).
Therefore, for normal Creo spool data the correct equation is m = lambda * L.
Diameter is retained for auditing, gauge fallback, and volumetric/custom modes.
If ``density_mode`` is set to ``volumetric``, m = rho * A * L is used instead.

Cosmetic tape and overbraid can be calculated from user-defined rules when Creo
XML does not carry sufficient physical data. Those assumptions remain visible
in the output rows and config file.
"""

from __future__ import annotations

import argparse
import dataclasses
import fnmatch
import json
import math
import re
import shutil
import sys
import time
import xml.etree.ElementTree as ET
from collections import defaultdict
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Iterable, Optional, Sequence

try:
    import creopyson
except Exception:  # optional unless --export-xml is used
    creopyson = None

try:
    from openpyxl import load_workbook
    from openpyxl.chart import BarChart, Reference
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
except ImportError as exc:
    raise SystemExit("openpyxl is required. Install requirements.txt") from exc

try:
    from tqdm import tqdm
except ImportError as exc:
    raise SystemExit("tqdm is required. Install requirements.txt") from exc


APP_VERSION = "1.0.0"

LENGTH_TO_M = {
    "m": 1.0,
    "meter": 1.0,
    "meters": 1.0,
    "mm": 1e-3,
    "millimeter": 1e-3,
    "millimeters": 1e-3,
    "cm": 1e-2,
    "centimeter": 1e-2,
    "centimeters": 1e-2,
    "in": 0.0254,
    "inch": 0.0254,
    "inches": 0.0254,
    "ft": 0.3048,
    "foot": 0.3048,
    "feet": 0.3048,
}

MASS_TO_KG = {
    "kg": 1.0,
    "kilogram": 1.0,
    "kilograms": 1.0,
    "g": 1e-3,
    "gram": 1e-3,
    "grams": 1e-3,
    "mg": 1e-6,
    "lb": 0.45359237,
    "lbs": 0.45359237,
    "lbm": 0.45359237,
    "oz": 0.028349523125,
}


def info(msg: str) -> None:
    tqdm.write(f"[INFO] {msg}")


def warn(msg: str) -> None:
    tqdm.write(f"[WARNING] {msg}")


def norm_key(value: str) -> str:
    return re.sub(r"[^A-Z0-9]+", "_", value.upper()).strip("_")


def clean_unit(value: Optional[str]) -> str:
    if not value:
        return ""
    return norm_key(str(value)).lower()


def as_float(value: Any) -> Optional[float]:
    if value is None:
        return None
    try:
        return float(str(value).strip())
    except (TypeError, ValueError):
        return None


def first_value(fields: dict[str, str], *names: str) -> Optional[str]:
    for name in names:
        key = norm_key(name)
        if key in fields and str(fields[key]).strip() != "":
            return str(fields[key]).strip()
    return None


def awg_to_diameter_m(gauge: Any) -> Optional[float]:
    """Bare conductor diameter from AWG. Returns metres."""
    if gauge is None:
        return None
    match = re.search(r"(?<!\d)(\d{1,2})(?:\s*AWG)?", str(gauge).upper())
    if not match:
        return None
    n = int(match.group(1))
    if not 0 <= n <= 40:
        return None
    d_mm = 0.127 * (92.0 ** ((36.0 - n) / 39.0))
    return d_mm / 1000.0


def length_to_m(value: Any, unit: Optional[str]) -> Optional[float]:
    val = as_float(value)
    if val is None:
        return None
    factor = LENGTH_TO_M.get(clean_unit(unit))
    if factor is None:
        return None
    return val * factor


def linear_density_to_kg_per_m(
    density: Any,
    mass_unit: Optional[str],
    length_unit: Optional[str],
) -> Optional[float]:
    value = as_float(density)
    if value is None:
        return None
    mf = MASS_TO_KG.get(clean_unit(mass_unit))
    lf = LENGTH_TO_M.get(clean_unit(length_unit))
    if mf is None or lf is None or lf == 0:
        return None
    return value * mf / lf


def volumetric_density_to_kg_m3(value: Any, unit: Optional[str]) -> Optional[float]:
    v = as_float(value)
    if v is None:
        return None
    u = clean_unit(unit)
    factors = {
        "kg_m3": 1.0,
        "kg_per_m3": 1.0,
        "g_cm3": 1000.0,
        "g_per_cm3": 1000.0,
        "lb_ft3": 16.01846337396,
        "lb_per_ft3": 16.01846337396,
    }
    if u in factors:
        return v * factors[u]
    return None


@dataclass
class SourceHarness:
    name: str
    indices: list[str] = field(default_factory=list)
    immediate_assembly: str = ""
    owner: str = ""
    created_by: str = ""
    last_modified_by: str = ""


@dataclass
class Spool:
    name: str
    kind: str = "WIRE"
    gauge: Optional[str] = None
    diameter_m: Optional[float] = None
    diameter_source: str = ""
    density_raw: Optional[float] = None
    density_mode: str = "linear"
    density_si: Optional[float] = None  # kg/m linear or kg/m3 volumetric
    mass_unit: Optional[str] = None
    length_unit: Optional[str] = None
    material: Optional[str] = None
    raw_fields: dict[str, str] = field(default_factory=dict)


@dataclass
class Connection:
    harness: str
    name: str
    spool_name: str
    length_m: Optional[float]
    length_raw: Optional[float] = None
    length_unit: Optional[str] = None
    from_ref: Optional[str] = None
    to_ref: Optional[str] = None
    raw_fields: dict[str, str] = field(default_factory=dict)


@dataclass
class MassRow:
    harness: str
    connection: str
    spool: str
    item_type: str
    gauge: Optional[str]
    diameter_m: Optional[float]
    density_mode: str
    density_si: Optional[float]
    length_m: Optional[float]
    mass_kg: Optional[float]
    from_ref: Optional[str] = None
    to_ref: Optional[str] = None
    material: Optional[str] = None
    assumption: str = ""
    xml_file: str = ""
    warning: str = ""


@dataclass
class CosmeticRule:
    name: str
    kind: str
    harness_pattern: str = "*"
    connection_pattern: str = "*"
    length_m: Optional[float] = None
    length_factor: float = 1.0
    thickness_mm: Optional[float] = None
    width_mm: Optional[float] = None
    material_density_kg_m3: Optional[float] = None
    areal_density_kg_m2: Optional[float] = None
    coverage_fraction: float = 1.0
    overlap_fraction: float = 0.5
    diameter_mm: Optional[float] = None


@dataclass
class Config:
    harness_tokens: list[str] = field(default_factory=lambda: ["_ROUT"])
    xml_dir: Path = Path("harness_xml")
    density_mode: str = "auto"
    default_length_unit: str = "mm"
    default_mass_unit: str = "kg"
    default_volumetric_density_unit: str = "kg/m3"
    mapkey_template: Optional[Path] = None
    mapkey_delay_ms: int = 1500
    creo_host: str = "localhost"
    creo_port: int = 9056
    cosmetic_rules: list[CosmeticRule] = field(default_factory=list)


def load_config(path: Optional[Path]) -> Config:
    cfg = Config()
    if path is None:
        return cfg
    data = json.loads(path.read_text(encoding="utf-8"))
    for key in (
        "harness_tokens",
        "density_mode",
        "default_length_unit",
        "default_mass_unit",
        "default_volumetric_density_unit",
        "mapkey_delay_ms",
        "creo_host",
        "creo_port",
    ):
        if key in data:
            setattr(cfg, key, data[key])
    if data.get("xml_dir"):
        cfg.xml_dir = Path(data["xml_dir"])
    if data.get("mapkey_template"):
        cfg.mapkey_template = Path(data["mapkey_template"])
    cfg.cosmetic_rules = [CosmeticRule(**x) for x in data.get("cosmetic_rules", [])]
    return cfg


def matches_harness(name: str, tokens: Sequence[str]) -> bool:
    upper = name.upper()
    return any(token.upper() in upper for token in tokens if token)


def discover_from_json(path: Path, tokens: Sequence[str]) -> list[SourceHarness]:
    data = json.loads(path.read_text(encoding="utf-8"))
    rows = data.get("hierarchy", data if isinstance(data, list) else [])
    found: dict[str, SourceHarness] = {}
    for row in rows:
        name = str(row.get("name") or row.get("Name") or "")
        if not name or not matches_harness(name, tokens):
            continue
        if not name.lower().endswith(".asm"):
            continue
        rec = found.setdefault(name, SourceHarness(name=name))
        indices = row.get("indices") or row.get("Index") or []
        if isinstance(indices, str):
            indices = [x.strip() for x in re.split(r"[\n;,]+", indices) if x.strip()]
        rec.indices.extend(x for x in indices if x not in rec.indices)
        rec.immediate_assembly = str(row.get("immediate_assembly") or "")
        rec.owner = str(row.get("owner") or "")
        rec.created_by = str(row.get("created_by") or "")
        rec.last_modified_by = str(row.get("last_modified_by") or "")
    return sorted(found.values(), key=lambda x: x.name.casefold())


def discover_from_xlsx(path: Path, tokens: Sequence[str]) -> list[SourceHarness]:
    wb = load_workbook(path, read_only=True, data_only=True)
    if "BOM Hierarchy" not in wb.sheetnames:
        raise ValueError("Workbook has no 'BOM Hierarchy' sheet.")
    ws = wb["BOM Hierarchy"]
    headers = {str(c.value).strip(): i + 1 for i, c in enumerate(ws[1]) if c.value}
    name_col = headers.get("Name")
    if not name_col:
        raise ValueError("BOM Hierarchy has no Name column.")
    found: dict[str, SourceHarness] = {}
    for r in range(2, ws.max_row + 1):
        name = str(ws.cell(r, name_col).value or "")
        if not name.lower().endswith(".asm") or not matches_harness(name, tokens):
            continue
        rec = found.setdefault(name, SourceHarness(name=name))
        idx = str(ws.cell(r, headers.get("Index", 1)).value or "")
        for x in re.split(r"[\n;,]+", idx):
            if x.strip() and x.strip() not in rec.indices:
                rec.indices.append(x.strip())
        for attr, header in (
            ("immediate_assembly", "Immediate Assembly"),
            ("owner", "Owner"),
            ("created_by", "Created By"),
            ("last_modified_by", "Last Modified By"),
        ):
            col = headers.get(header)
            if col:
                setattr(rec, attr, str(ws.cell(r, col).value or ""))
    wb.close()
    return sorted(found.values(), key=lambda x: x.name.casefold())


def discover_harnesses(source: Path, tokens: Sequence[str]) -> list[SourceHarness]:
    if source.suffix.lower() == ".json":
        return discover_from_json(source, tokens)
    if source.suffix.lower() in {".xlsx", ".xlsm"}:
        return discover_from_xlsx(source, tokens)
    raise ValueError("Source must be .json or .xlsx/.xlsm")


def safe_xml_name(model: str) -> str:
    stem = re.sub(r"\.(asm|prt)$", "", model, flags=re.I)
    return re.sub(r"[^A-Za-z0-9_.-]+", "_", stem) + ".xml"


def export_xml_with_mapkey(
    harnesses: Sequence[SourceHarness], cfg: Config, overwrite: bool = False
) -> None:
    if creopyson is None:
        raise RuntimeError("creopyson is required for --export-xml")
    if cfg.mapkey_template is None:
        raise ValueError("--export-xml requires mapkey_template in the config file.")
    template = cfg.mapkey_template.read_text(encoding="utf-8")
    if "{xml_path}" not in template:
        raise ValueError("Mapkey template must contain the placeholder {xml_path}.")
    cfg.xml_dir.mkdir(parents=True, exist_ok=True)

    client = creopyson.Client(cfg.creo_host, cfg.creo_port)
    client.connect()
    for harness in tqdm(harnesses, desc="Exporting harness XML", unit="harness"):
        xml_path = (cfg.xml_dir / safe_xml_name(harness.name)).resolve()
        if xml_path.exists() and not overwrite:
            continue
        client.file_display(harness.name, activate=True)
        script = template.replace("{xml_path}", str(xml_path).replace("\\", "/"))
        script = script.replace("{harness}", harness.name)
        client.interface_mapkey(script, delay=cfg.mapkey_delay_ms)
        # Mapkey call returning does not guarantee all Creo file I/O is flushed.
        for _ in range(20):
            if xml_path.exists():
                break
            time.sleep(0.25)
        if not xml_path.exists():
            warn(f"No XML appeared for {harness.name}: {xml_path}")


def element_fields(elem: ET.Element) -> dict[str, str]:
    fields: dict[str, str] = {}
    for k, v in elem.attrib.items():
        fields[norm_key(k)] = str(v).strip()
    if elem.text and elem.text.strip() and len(list(elem)) == 0:
        fields[norm_key(elem.tag)] = elem.text.strip()
    for child in elem:
        tag = norm_key(child.tag)
        if child.text and child.text.strip() and len(list(child)) == 0:
            fields.setdefault(tag, child.text.strip())
        for k, v in child.attrib.items():
            fields.setdefault(norm_key(f"{child.tag}_{k}"), str(v).strip())
            # XML schemas often encode parameter entries as name/value pairs.
            if norm_key(k) in {"NAME", "PARAMETER", "PARAM_NAME"}:
                val = child.attrib.get("value") or child.attrib.get("Value") or child.text
                if val:
                    fields[norm_key(str(v))] = str(val).strip()
    return fields


def iter_element_records(root: ET.Element) -> Iterable[tuple[str, dict[str, str]]]:
    for elem in root.iter():
        fields = element_fields(elem)
        if fields:
            yield norm_key(elem.tag), fields


def parse_spools(root: ET.Element, cfg: Config) -> dict[str, Spool]:
    candidates: dict[str, Spool] = {}
    for tag, fields in iter_element_records(root):
        name = first_value(fields, "SPOOL_NAME", "NAME", "SPOOL")
        has_spool_signal = "SPOOL" in tag or first_value(fields, "SPOOL_NAME") is not None
        has_physical = any(
            first_value(fields, key) is not None
            for key in ("DENSITY", "THICKNESS", "OUTER_DIAMETER", "WIRE_GAUGE", "GAUGE")
        )
        if not name or not has_spool_signal or not has_physical:
            continue

        length_unit = first_value(fields, "UNITS", "LENGTH_UNITS", "UNIT") or cfg.default_length_unit
        mass_unit = first_value(fields, "MASS_UNITS", "MASS_UNIT") or cfg.default_mass_unit
        gauge = first_value(fields, "WIRE_GAUGE", "GAUGE")
        diameter_raw = first_value(fields, "THICKNESS", "OUTER_DIAMETER", "DIAMETER")
        diameter_m = length_to_m(diameter_raw, length_unit) if diameter_raw else None
        diameter_source = "XML THICKNESS/DIAMETER" if diameter_m is not None else ""
        if diameter_m is None and gauge:
            diameter_m = awg_to_diameter_m(gauge)
            if diameter_m is not None:
                diameter_source = "AWG bare-conductor fallback"

        density_raw = as_float(first_value(fields, "DENSITY", "LINEAR_DENSITY"))
        mode = cfg.density_mode.lower()
        if mode == "auto":
            # PTC spool DENSITY is documented as linear density.
            mode = "linear"
        if mode == "linear":
            density_si = linear_density_to_kg_per_m(density_raw, mass_unit, length_unit)
        elif mode == "volumetric":
            density_unit = first_value(fields, "DENSITY_UNITS") or cfg.default_volumetric_density_unit
            density_si = volumetric_density_to_kg_m3(density_raw, density_unit)
        else:
            raise ValueError("density_mode must be auto, linear, or volumetric")

        kind = first_value(fields, "TYPE", "OBJ_TYPE") or "WIRE"
        material = first_value(fields, "MATERIAL", "INSUL_TYPE", "WIRE_CONSTRUCTION")
        spool = Spool(
            name=name,
            kind=kind,
            gauge=gauge,
            diameter_m=diameter_m,
            diameter_source=diameter_source,
            density_raw=density_raw,
            density_mode=mode,
            density_si=density_si,
            mass_unit=mass_unit,
            length_unit=length_unit,
            material=material,
            raw_fields=fields,
        )
        old = candidates.get(name)
        if old is None or sum(v is not None for v in (spool.density_si, spool.diameter_m, spool.gauge)) > sum(
            v is not None for v in (old.density_si, old.diameter_m, old.gauge)
        ):
            candidates[name] = spool
    return candidates


def parse_connections(root: ET.Element, harness: str, cfg: Config) -> list[Connection]:
    out: list[Connection] = []
    seen: set[tuple[str, str, str]] = set()
    for tag, fields in iter_element_records(root):
        spool = first_value(fields, "SPOOL_NAME", "SPOOL", "SPOOL_REF")
        length_raw_s = first_value(
            fields,
            "LENGTH",
            "WIRE_LENGTH",
            "CABLE_LENGTH",
            "ROUTED_LENGTH",
            "TOTAL_LENGTH",
        )
        # Connections must map to a spool and provide a routed length.
        if not spool or length_raw_s is None:
            continue
        length_raw = as_float(length_raw_s)
        if length_raw is None:
            continue
        unit = first_value(fields, "UNITS", "LENGTH_UNITS", "UNIT") or cfg.default_length_unit
        length_m = length_to_m(length_raw, unit)
        name = first_value(fields, "NAME", "WIRE_NAME", "CABLE_NAME", "CONNECTION_NAME", "ID")
        if not name:
            name = f"{tag}_{len(out)+1}"
        from_ref = first_value(fields, "FROM", "FROM_REF", "FROM_PIN", "START", "ENTRY_PORT")
        to_ref = first_value(fields, "TO", "TO_REF", "TO_PIN", "END", "EXIT_PORT")
        key = (name, spool, f"{length_raw}:{unit}")
        if key in seen:
            continue
        seen.add(key)
        out.append(
            Connection(
                harness=harness,
                name=name,
                spool_name=spool,
                length_m=length_m,
                length_raw=length_raw,
                length_unit=unit,
                from_ref=from_ref,
                to_ref=to_ref,
                raw_fields=fields,
            )
        )
    return out


def parse_harness_xml(path: Path, harness: str, cfg: Config) -> tuple[dict[str, Spool], list[Connection]]:
    root = ET.parse(path).getroot()
    return parse_spools(root, cfg), parse_connections(root, harness, cfg)


def spool_mass_kg(spool: Spool, length_m: Optional[float]) -> Optional[float]:
    if length_m is None or spool.density_si is None:
        return None
    if spool.density_mode == "linear":
        return spool.density_si * length_m
    if spool.density_mode == "volumetric":
        if spool.diameter_m is None:
            return None
        area = math.pi * spool.diameter_m**2 / 4.0
        return spool.density_si * area * length_m
    return None


def cosmetic_mass_kg(rule: CosmeticRule, diameter_m: float, length_m: float) -> Optional[float]:
    kind = rule.kind.lower()
    thickness_m = (rule.thickness_mm or 0.0) / 1000.0
    diameter = (rule.diameter_mm / 1000.0) if rule.diameter_mm else diameter_m
    length = rule.length_m if rule.length_m is not None else length_m * rule.length_factor
    coverage = max(0.0, min(rule.coverage_fraction, 1.0))

    if kind == "overbraid":
        surface_area = math.pi * diameter * length
        if rule.areal_density_kg_m2 is not None:
            return surface_area * coverage * rule.areal_density_kg_m2
        if rule.material_density_kg_m3 is not None and thickness_m > 0:
            return surface_area * thickness_m * coverage * rule.material_density_kg_m3
        return None

    if kind == "tape":
        if rule.width_mm is None or rule.width_mm <= 0:
            return None
        width_m = rule.width_mm / 1000.0
        pitch = width_m * max(1.0 - rule.overlap_fraction, 1e-6)
        # Helical tape centerline length needed to advance axially by 'length'.
        tape_length = length * math.sqrt(1.0 + (math.pi * diameter / pitch) ** 2)
        if rule.areal_density_kg_m2 is not None:
            return tape_length * width_m * rule.areal_density_kg_m2
        if rule.material_density_kg_m3 is not None and thickness_m > 0:
            return tape_length * width_m * thickness_m * rule.material_density_kg_m3
        return None

    return None


def rule_matches(rule: CosmeticRule, harness: str, connection: str) -> bool:
    return fnmatch.fnmatch(harness.upper(), rule.harness_pattern.upper()) and fnmatch.fnmatch(
        connection.upper(), rule.connection_pattern.upper()
    )


def analyze_harness(
    harness: SourceHarness,
    xml_path: Path,
    cfg: Config,
) -> tuple[list[MassRow], dict[str, Any]]:
    spools, connections = parse_harness_xml(xml_path, harness.name, cfg)
    rows: list[MassRow] = []
    warnings: list[str] = []

    for conn in connections:
        spool = spools.get(conn.spool_name)
        if spool is None:
            rows.append(
                MassRow(
                    harness=harness.name,
                    connection=conn.name,
                    spool=conn.spool_name,
                    item_type="WIRE/CABLE",
                    gauge=None,
                    diameter_m=None,
                    density_mode="",
                    density_si=None,
                    length_m=conn.length_m,
                    mass_kg=None,
                    from_ref=conn.from_ref,
                    to_ref=conn.to_ref,
                    xml_file=str(xml_path),
                    warning="SPOOL_NOT_FOUND",
                )
            )
            warnings.append(f"{conn.name}: spool {conn.spool_name} not found")
            continue

        mass = spool_mass_kg(spool, conn.length_m)
        row_warning = ""
        if conn.length_m is None:
            row_warning = "LENGTH_UNIT_UNRESOLVED"
        elif spool.density_si is None:
            row_warning = "DENSITY_UNRESOLVED"
        elif spool.density_mode == "volumetric" and spool.diameter_m is None:
            row_warning = "DIAMETER_UNRESOLVED"
        rows.append(
            MassRow(
                harness=harness.name,
                connection=conn.name,
                spool=spool.name,
                item_type=spool.kind,
                gauge=spool.gauge,
                diameter_m=spool.diameter_m,
                density_mode=spool.density_mode,
                density_si=spool.density_si,
                length_m=conn.length_m,
                mass_kg=mass,
                from_ref=conn.from_ref,
                to_ref=conn.to_ref,
                material=spool.material,
                assumption=spool.diameter_source,
                xml_file=str(xml_path),
                warning=row_warning,
            )
        )
        if row_warning:
            warnings.append(f"{conn.name}: {row_warning}")

        # Cosmetic mass rules are connection-scoped; a length override can turn
        # them into harness-level assumptions if desired.
        for rule in cfg.cosmetic_rules:
            if not rule_matches(rule, harness.name, conn.name):
                continue
            if conn.length_m is None:
                continue
            d = (rule.diameter_mm / 1000.0) if rule.diameter_mm else spool.diameter_m
            if d is None:
                rows.append(
                    MassRow(
                        harness=harness.name,
                        connection=conn.name,
                        spool=rule.name,
                        item_type=rule.kind.upper(),
                        gauge=None,
                        diameter_m=None,
                        density_mode="cosmetic",
                        density_si=rule.material_density_kg_m3 or rule.areal_density_kg_m2,
                        length_m=rule.length_m or conn.length_m,
                        mass_kg=None,
                        xml_file=str(xml_path),
                        warning="COSMETIC_DIAMETER_UNRESOLVED",
                        assumption=rule.name,
                    )
                )
                continue
            cmass = cosmetic_mass_kg(rule, d, conn.length_m)
            assumption = (
                f"rule={rule.name}; kind={rule.kind}; thickness_mm={rule.thickness_mm}; "
                f"width_mm={rule.width_mm}; coverage={rule.coverage_fraction}; "
                f"overlap={rule.overlap_fraction}; length_override_m={rule.length_m}; "
                f"length_factor={rule.length_factor}"
            )
            rows.append(
                MassRow(
                    harness=harness.name,
                    connection=conn.name,
                    spool=rule.name,
                    item_type=rule.kind.upper(),
                    gauge=None,
                    diameter_m=d,
                    density_mode="cosmetic",
                    density_si=rule.material_density_kg_m3 or rule.areal_density_kg_m2,
                    length_m=rule.length_m or conn.length_m * rule.length_factor,
                    mass_kg=cmass,
                    material=None,
                    assumption=assumption,
                    xml_file=str(xml_path),
                    warning="" if cmass is not None else "COSMETIC_PROPERTIES_INCOMPLETE",
                )
            )

    total = sum(r.mass_kg for r in rows if r.mass_kg is not None)
    unresolved = sum(1 for r in rows if r.mass_kg is None)
    summary = {
        "harness": harness.name,
        "indices": harness.indices,
        "immediate_assembly": harness.immediate_assembly,
        "owner": harness.owner,
        "created_by": harness.created_by,
        "last_modified_by": harness.last_modified_by,
        "xml_file": str(xml_path),
        "spool_count": len(spools),
        "connection_count": len(connections),
        "mass_kg": total,
        "unresolved_rows": unresolved,
        "warnings": warnings,
    }
    return rows, summary


def analyze_all(harnesses: Sequence[SourceHarness], cfg: Config) -> tuple[list[MassRow], list[dict[str, Any]]]:
    rows: list[MassRow] = []
    summaries: list[dict[str, Any]] = []
    for harness in tqdm(harnesses, desc="Analyzing harness XML", unit="harness"):
        xml_path = cfg.xml_dir / safe_xml_name(harness.name)
        if not xml_path.exists():
            warn(f"XML not found for {harness.name}: {xml_path}")
            summaries.append(
                {
                    "harness": harness.name,
                    "indices": harness.indices,
                    "immediate_assembly": harness.immediate_assembly,
                    "owner": harness.owner,
                    "created_by": harness.created_by,
                    "last_modified_by": harness.last_modified_by,
                    "xml_file": str(xml_path),
                    "spool_count": 0,
                    "connection_count": 0,
                    "mass_kg": None,
                    "unresolved_rows": 1,
                    "warnings": ["XML_NOT_FOUND"],
                }
            )
            continue
        r, s = analyze_harness(harness, xml_path, cfg)
        rows.extend(r)
        summaries.append(s)
    return rows, summaries


def replace_sheet(wb, name: str):
    if name in wb.sheetnames:
        idx = wb.sheetnames.index(name)
        wb.remove(wb[name])
        return wb.create_sheet(name, idx)
    return wb.create_sheet(name)


def append_excel_sheets(workbook_path: Path, rows: Sequence[MassRow], summaries: Sequence[dict[str, Any]]) -> None:
    wb = load_workbook(workbook_path)
    ws = replace_sheet(wb, "Harness Masses")
    headers = [
        "Index",
        "Harness",
        "Immediate Assembly",
        "Owner",
        "Created By",
        "Last Modified By",
        "Connection",
        "Spool / Covering",
        "Item Type",
        "Gauge",
        "Diameter [mm]",
        "Density Mode",
        "Density [kg/m or kg/m^3]",
        "Length [m]",
        "Mass [kg]",
        "From",
        "To",
        "Material",
        "Assumption / Diameter Source",
        "XML File",
        "Warning",
    ]
    ws.append(headers)
    for cell in ws[1]:
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    summary_by_harness = {s["harness"]: s for s in summaries}
    for row in rows:
        meta = summary_by_harness.get(row.harness, {})
        ws.append(
            [
                "\n".join(meta.get("indices") or []),
                row.harness,
                meta.get("immediate_assembly"),
                meta.get("owner"),
                meta.get("created_by"),
                meta.get("last_modified_by"),
                row.connection,
                row.spool,
                row.item_type,
                row.gauge,
                row.diameter_m * 1000.0 if row.diameter_m is not None else None,
                row.density_mode,
                row.density_si,
                row.length_m,
                row.mass_kg,
                row.from_ref,
                row.to_ref,
                row.material,
                row.assumption,
                row.xml_file,
                row.warning,
            ]
        )
        rr = ws.max_row
        if row.warning:
            ws.cell(rr, 21).fill = PatternFill("solid", fgColor="FFF2CC")
        if row.item_type.upper() in {"TAPE", "OVERBRAID"}:
            for c in range(1, len(headers) + 1):
                ws.cell(rr, c).fill = PatternFill("solid", fgColor="E2F0D9")
        for c in range(1, len(headers) + 1):
            ws.cell(rr, c).alignment = Alignment(vertical="top", wrap_text=True)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    widths = [22, 28, 28, 20, 22, 24, 28, 24, 16, 12, 14, 16, 22, 14, 14, 22, 22, 20, 52, 50, 30]
    for i, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = width
    for r in range(2, ws.max_row + 1):
        for c in (11, 13, 14, 15):
            ws.cell(r, c).number_format = "0.000000"

    ss = replace_sheet(wb, "Harness Summary")
    sheaders = [
        "Harness",
        "Indices",
        "Immediate Assembly",
        "Owner",
        "Created By",
        "Last Modified By",
        "XML File",
        "Spools",
        "Connections",
        "Calculated Mass [kg]",
        "Unresolved Rows",
        "Warnings",
    ]
    ss.append(sheaders)
    for cell in ss[1]:
        cell.fill = PatternFill("solid", fgColor="7030A0")
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for s in summaries:
        ss.append(
            [
                s["harness"],
                "\n".join(s.get("indices") or []),
                s.get("immediate_assembly"),
                s.get("owner"),
                s.get("created_by"),
                s.get("last_modified_by"),
                s.get("xml_file"),
                s.get("spool_count"),
                s.get("connection_count"),
                s.get("mass_kg"),
                s.get("unresolved_rows"),
                "; ".join(s.get("warnings") or []),
            ]
        )
    ss.freeze_panes = "A2"
    ss.auto_filter.ref = ss.dimensions
    for i, width in enumerate([30, 24, 28, 20, 22, 24, 50, 10, 12, 22, 16, 50], 1):
        ss.column_dimensions[get_column_letter(i)].width = width
    for r in range(2, ss.max_row + 1):
        ss.cell(r, 10).number_format = "0.000000"
        for c in range(1, len(sheaders) + 1):
            ss.cell(r, c).alignment = Alignment(vertical="top", wrap_text=True)

    valid = [i for i, s in enumerate(summaries, start=2) if s.get("mass_kg") is not None]
    if valid:
        chart = BarChart()
        chart.title = "Calculated Wiring Harness Mass"
        chart.y_axis.title = "Mass [kg]"
        data = Reference(ss, min_col=10, min_row=min(valid), max_row=max(valid))
        cats = Reference(ss, min_col=1, min_row=min(valid), max_row=max(valid))
        chart.add_data(data, titles_from_data=False)
        chart.set_categories(cats)
        chart.height = 8
        chart.width = 14
        ss.add_chart(chart, "M2")

    wb.save(workbook_path)


def write_json_output(path: Path, rows: Sequence[MassRow], summaries: Sequence[dict[str, Any]]) -> None:
    payload = {
        "schema": "creo-harness-mass-report/v1",
        "summaries": list(summaries),
        "rows": [dataclasses.asdict(r) for r in rows],
    }
    path.write_text(json.dumps(payload, indent=2), encoding="utf-8")


def parser() -> argparse.ArgumentParser:
    p = argparse.ArgumentParser(
        description="Discover Creo routing assemblies, parse Cabling XML, and calculate harness masses.",
        formatter_class=argparse.ArgumentDefaultsHelpFormatter,
    )
    p.add_argument("source", type=Path, help="Assembly report .json or .xlsx")
    p.add_argument("--config", type=Path, help="Harness analysis JSON config")
    p.add_argument("--token", action="append", help="Harness name token. Repeat as needed; overrides config.")
    p.add_argument("--xml-dir", type=Path, help="Folder containing/generated harness XML files")
    p.add_argument("--export-xml", action="store_true", help="Export XML from Creo before parsing")
    p.add_argument("--overwrite-xml", action="store_true", help="Re-export existing XML files")
    p.add_argument("--excel", type=Path, help="Workbook to append Harness Masses/Summary sheets to")
    p.add_argument("--json-output", type=Path, help="Harness analysis JSON output")
    return p


def main(argv: Optional[Sequence[str]] = None) -> int:
    args = parser().parse_args(argv)
    cfg = load_config(args.config)
    if args.token:
        cfg.harness_tokens = args.token
    if args.xml_dir:
        cfg.xml_dir = args.xml_dir
    cfg.xml_dir = cfg.xml_dir.resolve()

    harnesses = discover_harnesses(args.source.resolve(), cfg.harness_tokens)
    info(f"Found {len(harnesses)} harness routing assembly file(s): {', '.join(h.name for h in harnesses) or '<none>'}")
    if not harnesses:
        return 0

    if args.export_xml:
        export_xml_with_mapkey(harnesses, cfg, overwrite=args.overwrite_xml)

    rows, summaries = analyze_all(harnesses, cfg)
    info(f"Generated {len(rows)} detailed harness mass row(s).")

    excel_path = args.excel
    if excel_path is None and args.source.suffix.lower() in {".xlsx", ".xlsm"}:
        excel_path = args.source
    if excel_path:
        append_excel_sheets(excel_path.resolve(), rows, summaries)
        info(f"Updated workbook: {excel_path.resolve()}")

    json_out = args.json_output or args.source.with_name(args.source.stem + "_harness.json")
    write_json_output(json_out.resolve(), rows, summaries)
    info(f"Harness JSON: {json_out.resolve()}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
